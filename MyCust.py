"""SQLite 기반 고객 정보 관리 프로그램."""

from __future__ import annotations

import sqlite3
import sys
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from MyCustManager import CustomerManager


def get_database_path() -> Path:
    """실행 방식에 맞는 영속 가능한 데이터베이스 절대 경로를 반환한다."""
    if getattr(sys, "frozen", False):
        executable_path = Path(sys.executable).resolve()

        # macOS PyInstaller .app의 Contents/MacOS/<executable>에서 .app 루트를 찾는다.
        for directory in (executable_path.parent, *executable_path.parents):
            if directory.suffix.lower() == ".app":
                return directory / "MyCust.db"

        # .app이 아닌 one-folder/one-file 패키지는 실행 파일 옆에 저장한다.
        return executable_path.parent / "MyCust.db"

    return Path(__file__).resolve().parent / "MyCust.db"


def export_customers_to_excel(
    customers: list[sqlite3.Row], output_path: Path
) -> None:
    """고객 전체 목록을 openpyxl 형식의 Excel 파일로 저장한다."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "고객 목록"
    worksheet.append(["고객 ID", "고객 이름", "고객 직함"])

    for customer in customers:
        worksheet.append(
            [customer["custID"], customer["custName"], customer["custTitle"]]
        )

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, width in enumerate((12, 24, 24), start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    workbook.save(output_path)


class CustomerView(QMainWindow):
    """고객 관리 화면 및 사용자 상호작용을 담당한다."""

    def __init__(self, customer_manager: CustomerManager) -> None:
        super().__init__()
        self.customer_manager = customer_manager
        self.setup_ui()
        self.connect_signals()
        self.load_customers()

    def setup_ui(self) -> None:
        self.setWindowTitle("고객 정보 관리")
        self.resize(760, 520)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)

        top_layout = QHBoxLayout()
        button_layout = QVBoxLayout()
        self.create_button = QPushButton("입력")
        self.update_button = QPushButton("수정")
        self.delete_button = QPushButton("삭제")
        self.search_button = QPushButton("검색")
        self.excel_save_button = QPushButton("엑셀저장")
        for button in (
            self.create_button,
            self.update_button,
            self.delete_button,
            self.search_button,
            self.excel_save_button,
        ):
            button.setMinimumWidth(100)
            button_layout.addWidget(button)
        button_layout.addStretch()

        form_layout = QFormLayout()
        self.id_edit = QLineEdit()
        self.id_edit.setReadOnly(True)
        self.name_edit = QLineEdit()
        self.title_edit = QLineEdit()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("이름 또는 직함으로 검색")
        form_layout.addRow("고객 ID", self.id_edit)
        form_layout.addRow("고객 이름", self.name_edit)
        form_layout.addRow("고객 직함", self.title_edit)
        form_layout.addRow("검색어", self.search_edit)

        top_layout.addLayout(button_layout)
        top_layout.addLayout(form_layout, 1)
        main_layout.addLayout(top_layout)

        self.customer_table = QTableWidget(0, 3)
        self.customer_table.setHorizontalHeaderLabels(["고객 ID", "고객 이름", "고객 직함"])
        self.customer_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.customer_table.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        self.customer_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.customer_table.horizontalHeader().setStretchLastSection(True)
        self.customer_table.setColumnWidth(0, 100)
        self.customer_table.setColumnWidth(1, 220)
        main_layout.addWidget(self.customer_table, 1)

    def connect_signals(self) -> None:
        self.create_button.clicked.connect(self.handle_create)
        self.update_button.clicked.connect(self.handle_update)
        self.delete_button.clicked.connect(self.handle_delete)
        self.search_button.clicked.connect(self.handle_search)
        self.excel_save_button.clicked.connect(self.handle_export_excel)
        self.search_edit.returnPressed.connect(self.handle_search)
        self.customer_table.itemSelectionChanged.connect(self.populate_form_from_selection)

    def load_customers(self, keyword: Optional[str] = None) -> None:
        try:
            customers = self.customer_manager.get_customers(keyword)
        except sqlite3.Error as error:
            self.show_database_error(error)
            return

        self.customer_table.setUpdatesEnabled(False)
        self.customer_table.clearContents()
        self.customer_table.setRowCount(len(customers))
        for row_index, customer in enumerate(customers):
            for column_index, column_name in enumerate(
                ("custID", "custName", "custTitle")
            ):
                item = QTableWidgetItem(str(customer[column_name]))
                if column_index == 0:
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.customer_table.setItem(row_index, column_index, item)
        self.customer_table.setUpdatesEnabled(True)

    def handle_create(self) -> None:
        customer_data = self.validated_form_data()
        if customer_data is None:
            return

        try:
            self.customer_manager.create_customer(*customer_data)
        except sqlite3.Error as error:
            self.show_database_error(error)
            return

        self.load_customers(self.search_edit.text())
        self.clear_form(keep_search=True)

    def handle_update(self) -> None:
        cust_id = self.selected_customer_id()
        if cust_id is None:
            self.show_warning("수정할 고객을 목록에서 선택하세요.")
            return
        customer_data = self.validated_form_data()
        if customer_data is None:
            return

        try:
            self.customer_manager.update_customer(cust_id, *customer_data)
        except (sqlite3.Error, ValueError) as error:
            self.show_database_error(error)
            return

        self.load_customers(self.search_edit.text())
        self.clear_form(keep_search=True)

    def handle_delete(self) -> None:
        cust_id = self.selected_customer_id()
        if cust_id is None:
            self.show_warning("삭제할 고객을 목록에서 선택하세요.")
            return

        answer = QMessageBox.question(
            self,
            "고객 삭제",
            "선택한 고객 정보를 삭제할까요?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        try:
            self.customer_manager.delete_customer(cust_id)
        except (sqlite3.Error, ValueError) as error:
            self.show_database_error(error)
            return

        self.load_customers(self.search_edit.text())
        self.clear_form(keep_search=True)

    def handle_search(self) -> None:
        self.load_customers(self.search_edit.text())
        self.clear_form(keep_search=True)

    def handle_export_excel(self) -> None:
        try:
            customers = self.customer_manager.get_customers()
        except sqlite3.Error as error:
            self.show_database_error(error)
            return

        if not customers:
            self.show_warning("엑셀로 저장할 고객 정보가 없습니다.")
            return

        file_name, _ = QFileDialog.getSaveFileName(
            self,
            "엑셀 파일 저장",
            "Customers.xlsx",
            "Excel 파일 (*.xlsx)",
        )
        if not file_name:
            return

        output_path = Path(file_name)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")

        try:
            export_customers_to_excel(customers, output_path)
        except (OSError, ValueError) as error:
            QMessageBox.critical(self, "엑셀 저장 오류", str(error))
            return

        QMessageBox.information(
            self,
            "엑셀 저장 완료",
            f"고객 정보 {len(customers)}건을 저장했습니다.\n{output_path}",
        )

    def populate_form_from_selection(self) -> None:
        selected_items = self.customer_table.selectedItems()
        if not selected_items:
            return
        row = selected_items[0].row()
        self.id_edit.setText(self.customer_table.item(row, 0).text())
        self.name_edit.setText(self.customer_table.item(row, 1).text())
        self.title_edit.setText(self.customer_table.item(row, 2).text())

    def selected_customer_id(self) -> Optional[int]:
        selected_items = self.customer_table.selectedItems()
        if not selected_items:
            return None
        return int(self.customer_table.item(selected_items[0].row(), 0).text())

    def validated_form_data(self) -> Optional[tuple[str, str]]:
        name = self.name_edit.text().strip()
        title = self.title_edit.text().strip()
        if not name or not title:
            self.show_warning("고객 이름과 고객 직함을 모두 입력하세요.")
            return None
        return name, title

    def clear_form(self, keep_search: bool = False) -> None:
        self.customer_table.clearSelection()
        self.id_edit.clear()
        self.name_edit.clear()
        self.title_edit.clear()
        if not keep_search:
            self.search_edit.clear()

    def show_warning(self, message: str) -> None:
        QMessageBox.warning(self, "입력 확인", message)

    def show_database_error(self, error: Exception) -> None:
        QMessageBox.critical(self, "데이터베이스 오류", str(error))


def main() -> int:
    app = QApplication(sys.argv)
    try:
        customer_manager = CustomerManager(get_database_path())
    except sqlite3.Error as error:
        QMessageBox.critical(None, "데이터베이스 오류", str(error))
        return 1

    window = CustomerView(customer_manager)
    window.show()
    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
