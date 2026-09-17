from urllib.parse import parse_qs, urljoin, urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment


SEARCH_URL = (
    "https://search.naver.com/search.naver?"
    "ie=UTF-8&query=%EB%B0%98%EB%8F%84%EC%B2%B4&sm=chr_hty"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    )
}
TIMEOUT = 15


def get_soup(url):
    response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    return BeautifulSoup(response.content, "html.parser")


def get_news_links(search_url, limit=10):
    soup = get_soup(search_url)
    news_links = []
    seen = set()

    for anchor in soup.select("a[href]"):
        href = anchor.get("href")
        if not isinstance(href, str):
            continue

        link = urljoin(search_url, href)
        parsed_link = urlparse(link)

        if parsed_link.netloc not in {"news.naver.com", "n.news.naver.com"}:
            continue
        if not parsed_link.path.startswith("/mnews/article/"):
            continue
        if link in seen:
            continue

        seen.add(link)
        news_links.append(link)
        if len(news_links) == limit:
            break

    return news_links


def clean_text(element):
    if element is None:
        return ""

    for unwanted in element.select("script, style, figure, .byline, .copyright"):
        unwanted.decompose()

    return " ".join(element.stripped_strings)


def get_article(article_url):
    soup = get_soup(article_url)

    title = soup.select_one(
        "h2#title_area, h2#newsct_article_title, "
        "div.media_end_head_headline h2, meta[property='og:title']"
    )
    body = soup.select_one(
        "#dic_area, #newsct_article, #articleBodyContents, #articeBody"
    )
    date = soup.select_one(
        "._ARTICLE_DATE_TIME, time, span.t11, .media_end_head_info_datestamp_time"
    )

    title_text = ""
    if title is not None:
        title_text = title.get("content", "") if title.name == "meta" else clean_text(title)

    return {
        "title": title_text,
        "date": clean_text(date),
        "url": article_url,
        "content": clean_text(body),
    }


def save_to_excel(articles, output_file="naver_result.xlsx"):
    workbook = openpyxl.Workbook()
    worksheet = workbook.create_sheet("네이버뉴스")
    del workbook["Sheet"]

    headers = ["번호", "제목", "날짜", "URL", "본문"]
    worksheet.append(headers)

    for number, article in enumerate(articles, start=1):
        worksheet.append(
            [
                number,
                article["title"],
                article["date"],
                article["url"],
                article["content"],
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 45
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 60
    worksheet.column_dimensions["E"].width = 100

    for row in worksheet.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_file)


def main():
    query = parse_qs(urlparse(SEARCH_URL).query).get("query", ["반도체"])[0]
    article_urls = get_news_links(SEARCH_URL, limit=10)
    articles = []

    for article_url in article_urls:
        try:
            article = get_article(article_url)
        except requests.RequestException as error:
            print(f"기사 요청 실패: {article_url} ({error})")
            continue

        if not article["content"]:
            print(f"본문을 찾지 못했습니다: {article_url}")
            continue

        articles.append(article)
        print(f"수집 완료: {article['title']}")

    output_file = "naver_result.xlsx"
    save_to_excel(articles, output_file)
    print(f"총 {len(articles)}건 저장 완료: {output_file}")


if __name__ == "__main__":
    main()